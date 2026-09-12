import openpyxl
file = r'C:\Users\ADMIN\Desktop\关键词4-16.xlsx'
wb = openpyxl.load_workbook(file, data_only=True)
print("Sheets:", wb.sheetnames)

for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n--- Sheet: {name} ---")
    for row in ws.iter_rows(values_only=True):
        if any(cell is not None for cell in row):
            print(row)