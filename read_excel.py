import openpyxl
wb = openpyxl.load_workbook(r'C:\Users\ADMIN\Desktop\关键词4-16.xlsx')
for s in wb.sheetnames:
    ws = wb[s]
    print(f'=== Sheet: {s} ===')
    for row in ws.iter_rows(max_row=15, values_only=True):
        row_data = [c for c in row if c is not None]
        if row_data:
            print(row_data)